import pptx

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx")
print(f"Template total slides: {len(prs.slides)}")
for idx, slide in enumerate(prs.slides):
    print(f"\n=================== TEMPLATE SLIDE {idx+1} ===================")
    print(f"Slide shapes count: {len(slide.shapes)}")
    for s_idx, shape in enumerate(slide.shapes):
        txt = ""
        if shape.has_text_frame:
            txt = " | ".join([p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
        pic_info = ""
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            pic_info = f" [PIC: {shape.image.filename}]"
        print(f"Shape {s_idx:02d} ({shape.name}) type={shape.shape_type} pos=({shape.left.inches:.2f}, {shape.top.inches:.2f}, {shape.width.inches:.2f}, {shape.height.inches:.2f}){pic_info} : {txt[:80]}")
