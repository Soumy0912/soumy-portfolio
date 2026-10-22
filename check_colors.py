import pptx

def inspect_colors_fonts(path, name):
    prs = pptx.Presentation(path)
    print(f"=== {name} ===")
    for idx, slide in enumerate(prs.slides):
        if idx >= 3:
            break
        print(f"Slide {idx+1}:")
        for s in slide.shapes:
            if s.has_text_frame:
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        font_name = r.font.name
                        font_size = r.font.size.pt if r.font.size else None
                        color = r.font.color.rgb if r.font.color and r.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB else None
                        if r.text.strip():
                            print(f"   text: '{r.text.strip()[:30]}' font={font_name} size={font_size} color={color}")
                            break

inspect_colors_fonts(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx", "TEMPLATE")
print("\n" + "="*50 + "\n")
inspect_colors_fonts(r"C:\Users\SHREE RAM\Downloads\SOHAM DALAL PORTFOLIO 1.pptx", "SOHAM DALAL")
