import pptx

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx")
print("=================== TEMPLATE SLIDES ===================")
for idx, slide in enumerate(prs.slides):
    print(f"--- SLIDE {idx+1} ---")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                t = p.text.strip()
                if t:
                    print(f"  [shape {s_idx} p {p_idx}] '{t}'")
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"  [PIC shape {s_idx}] image: {shape.image.filename} ({shape.width.inches:.2f} x {shape.height.inches:.2f} at {shape.left.inches:.2f}, {shape.top.inches:.2f})")
