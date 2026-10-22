import pptx

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx")
print("Listing all picture parts across slides:")
for idx, slide in enumerate(prs.slides):
    print(f"\nSlide {idx+1}:")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for sub_idx, sub in enumerate(shape.shapes):
                if sub.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                    rId = sub._element.blipFill.blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    part = slide.part.related_part(rId)
                    print(f"  Shape {s_idx}.{sub_idx} ({sub.name}): rId={rId}, partname={part.partname}, size={len(part.blob)}")
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            rId = shape._element.blipFill.blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            part = slide.part.related_part(rId)
            print(f"  Shape {s_idx} ({shape.name}): rId={rId}, partname={part.partname}, size={len(part.blob)}")
