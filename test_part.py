import pptx

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx")
slide2 = prs.slides[1]
group = slide2.shapes[4]
pic = group.shapes[0]
rId = pic._element.blipFill.blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
print("rId:", rId)
part = slide2.part.related_part(rId)
print("Related part:", type(part), part.partname)
print("Part blob length:", len(part.blob))
