import pptx

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx")
slide2 = prs.slides[1]
group = slide2.shapes[4]
pic = group.shapes[0]
rId = pic._element.blipFill.blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
part = slide2.part.related_part(rId)

# Read Soumy portrait image
with open(r"C:\Users\SHREE RAM\Downloads\Image (5).jpg", "rb") as f:
    new_data = f.read()

part._blob = new_data
test_out = r"C:\Users\SHREE RAM\.gemini\antigravity\scratch\soumy-portfolio\test_out.pptx"
prs.save(test_out)
print("Saved successfully!")

# Reopen and check
prs2 = pptx.Presentation(test_out)
slide2_new = prs2.slides[1]
group_new = slide2_new.shapes[4]
pic_new = group_new.shapes[0]
rId_new = pic_new._element.blipFill.blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
part_new = slide2_new.part.related_part(rId_new)
print("Reopened blob length:", len(part_new.blob), "matches new data:", len(part_new.blob) == len(new_data))
