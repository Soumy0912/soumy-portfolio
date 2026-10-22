import pptx

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\SOHAM DALAL PORTFOLIO 1.pptx")
print("=================== SOHAM DALAL FULL CONTENT ===================")
for idx, slide in enumerate(prs.slides):
    print(f"\n=================== SLIDE {idx+1} ===================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            full_txt = "\n".join([p.text for p in shape.text_frame.paragraphs if p.text.strip()])
            if full_txt.strip():
                print(f"--- Shape {s_idx} ({shape.name}) [{shape.left.inches:.2f}, {shape.top.inches:.2f}, {shape.width.inches:.2f}, {shape.height.inches:.2f}]:\n{full_txt}")
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"--- Picture Shape {s_idx} ({shape.name}) [{shape.left.inches:.2f}, {shape.top.inches:.2f}, {shape.width.inches:.2f}, {shape.height.inches:.2f}]: filename={shape.image.filename}")
