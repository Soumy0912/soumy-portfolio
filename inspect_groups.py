import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx")

for idx, slide in enumerate(prs.slides):
    print(f"\n=================== SLIDE {idx+1} GROUPS ===================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            print(f"Group {s_idx} ({shape.name}) with {len(shape.shapes)} sub-shapes:")
            for sub_i, sub_s in enumerate(shape.shapes):
                txt = ""
                if sub_s.has_text_frame:
                    txt = " | ".join([p.text.strip() for p in sub_s.text_frame.paragraphs if p.text.strip()])
                pic = f"PIC({sub_s.image.filename})" if sub_s.shape_type == MSO_SHAPE_TYPE.PICTURE else ""
                print(f"   Sub {sub_i:02d}: {sub_s.name} type={sub_s.shape_type} {pic} pos=({sub_s.left.inches:.2f}, {sub_s.top.inches:.2f}) {txt}")
