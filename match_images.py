import pptx
import hashlib
import os

def file_hash(path):
    with open(path, 'rb') as f:
        return hashlib.md5(f.read()).hexdigest()

print("Hashing downloads images:")
dl_hashes = {}
dl_dir = r"C:\Users\SHREE RAM\Downloads"
for f in os.listdir(dl_dir):
    if f.lower().endswith(('.jpg', '.jpeg', '.png')):
        p = os.path.join(dl_dir, f)
        h = file_hash(p)
        dl_hashes[h] = f

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\SOHAM DALAL PORTFOLIO 1.pptx")
print("\nChecking Soham Dalal slides pictures:")
for idx, slide in enumerate(prs.slides):
    for s in slide.shapes:
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            img_bytes = s.image.blob
            h = hashlib.md5(img_bytes).hexdigest()
            match = dl_hashes.get(h, "NO MATCH IN DL")
            print(f"Slide {idx+1}: {s.name} len={len(img_bytes)} -> matched {match}")

print("\nChecking template pictures:")
prs_tpl = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx")
for idx, slide in enumerate(prs_tpl.slides):
    for s in slide.shapes:
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for sub in s.shapes:
                if sub.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                    img_bytes = sub.image.blob
                    h = hashlib.md5(img_bytes).hexdigest()
                    match = dl_hashes.get(h, "NO MATCH IN DL")
                    print(f"Template Slide {idx+1}: {sub.name} len={len(img_bytes)} -> matched {match}")
