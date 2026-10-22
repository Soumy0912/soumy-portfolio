import pptx

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx")
slide2 = prs.slides[1]
group = slide2.shapes[4]
pic = group.shapes[0]
print("Pic shape:", pic.name, "image rId:", pic._element.blipFill.blip.attrib)
image_part = slide2.part.related_slide(pic._element.blipFill.blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')) if hasattr(slide2.part, 'related_slide') else None
print("Image part:", pic.image.part)
print("Blob size:", len(pic.image.blob))
# Can we update the blob?
# pic.image._blob = new_bytes
