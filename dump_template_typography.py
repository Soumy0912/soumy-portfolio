import pptx

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\Green and Yellow Bold Playful Creative Portfolio Presentation.pptx")

for idx, slide in enumerate(prs.slides):
    print(f"\n=======================================================")
    print(f"SLIDE {idx+1}")
    print(f"=======================================================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            tf = shape.text_frame
            print(f"Shape {s_idx:02d} ({shape.name}): [x={shape.left.inches:.2f}, y={shape.top.inches:.2f}, w={shape.width.inches:.2f}, h={shape.height.inches:.2f}]")
            for p_idx, p in enumerate(tf.paragraphs):
                p_text = p.text.strip()
                font_info = ""
                if p.runs:
                    r = p.runs[0]
                    col = r.font.color.rgb if r.font.color and r.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB else None
                    font_info = f"font={r.font.name}, sz={r.font.size.pt if r.font.size else None}, col={col}, bold={r.font.bold}"
                print(f"   P{p_idx:02d} [{font_info}]: '{p_text}'")
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            print(f"Shape {s_idx:02d} (GROUP: {shape.name}) [x={shape.left.inches:.2f}, y={shape.top.inches:.2f}, w={shape.width.inches:.2f}, h={shape.height.inches:.2f}] with {len(shape.shapes)} subshapes")
            for sub_i, sub in enumerate(shape.shapes):
                if sub.has_text_frame:
                    for p_i, p in enumerate(sub.text_frame.paragraphs):
                        print(f"      SubShape {sub_i} text: '{p.text.strip()}'")
                elif sub.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                    print(f"      SubShape {sub_i} PICTURE: {sub.name}")
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"Shape {s_idx:02d} (PICTURE: {shape.name}) [x={shape.left.inches:.2f}, y={shape.top.inches:.2f}, w={shape.width.inches:.2f}, h={shape.height.inches:.2f}]")
