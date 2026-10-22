import pptx

prs = pptx.Presentation(r"C:\Users\SHREE RAM\Downloads\SOUMY RAGHUVANSHI PORTFOLIO PRESENTATION.pptx")
print(f"Total slides: {len(prs.slides)}")

lorem_count = 0
margarita_count = 0

for idx, slide in enumerate(prs.slides):
    print(f"\n=================== SLIDE {idx+1} ===================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                t = p.text.strip()
                if t:
                    if "lorem" in t.lower():
                        lorem_count += 1
                        print(f"  [LOREM WARNING] Shape {s_idx}: {t}")
                    if "margarita" in t.lower():
                        margarita_count += 1
                        print(f"  [MARGARITA WARNING] Shape {s_idx}: {t}")
                    r = p.runs[0] if p.runs else None
                    font_name = r.font.name if r else "Default"
                    size_pt = r.font.size.pt if r and r.font.size else "Auto"
                    print(f"  Shape {s_idx} [{font_name}, {size_pt}pt]: {t[:85]}")
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for sub_i, sub in enumerate(shape.shapes):
                if sub.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                    rId = sub._element.blipFill.blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    part = slide.part.related_part(rId)
                    print(f"  Group {s_idx}.{sub_i} PIC: {part.partname} ({len(part.blob)} bytes)")

print(f"\nAudit complete: Lorem occurrences = {lorem_count}, Margarita occurrences = {margarita_count}")
